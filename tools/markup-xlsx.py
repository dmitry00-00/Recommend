# -*- coding: utf-8 -*-
"""Собрать film_reviews.xlsx из .cache/markup/film-reviews.json (его пишет markup-xlsx.mts).

    python3 tools/markup-xlsx.py [--empty 300] [--out .cache/markup/film_reviews.xlsx]

Нужен openpyxl (pip install openpyxl) — только ради выпадающего списка: он живёт в самом
файле xlsx, CSV его не умеет. Список фильмов лежит на скрытом листе «Справочник» и подключён
к колонкам «Фильм» именованным диапазоном film_list; там же, во второй колонке, ключ фильма —
по нему разметку читают обратно.

Листов с разметкой два — «Обзоры» и «Эссе» (просьба владельца 26.09). Это разная работа:
у обзорщика фильм почти всегда в заголовке ролика, у эссеиста заголовок бывает про что угодно.
Третий лист, «Без разбора», — обратный взгляд: фильмы, у которых нет ни обзора, ни разбора;
там же можно вписать ссылку, если разбор всё-таки знаешь.
"""
import json, sys, os
from openpyxl import Workbook
from openpyxl.styles import Font, Alignment, PatternFill
from openpyxl.worksheet.datavalidation import DataValidation
from openpyxl.workbook.defined_name import DefinedName
from openpyxl.utils import get_column_letter

def arg(name, default):
    return sys.argv[sys.argv.index('--' + name) + 1] if '--' + name in sys.argv[:-1] else default

ROOT = os.path.dirname(os.path.dirname(os.path.abspath(__file__)))
EMPTY = int(arg('empty', '300'))
OUT = arg('out', os.path.join(ROOT, '.cache/markup/film_reviews.xlsx'))
NOT_A_FILM = '— не про фильм —'
HEAD_FILL = PatternFill('solid', fgColor='EFEFEF')
LINK = Font(color='0563C1', underline='single')

data = json.load(open(os.path.join(ROOT, '.cache/markup/film-reviews.json'), encoding='utf-8'))
films, missing = data['films'], data['missing']

wb = Workbook()

# 1. Служебный лист: сначала «не про фильм» (у обзорщиков много роликов не о кино), дальше
#    все фильмы по алфавиту. Ключ рядом — чтобы заполненную таблицу можно было прочитать.
ref = wb.create_sheet('Справочник')
ref.append(['Фильм', 'ключ'])
ref.append([NOT_A_FILM, ''])
for f in films:
    ref.append([f['label'], f['key']])
ref.column_dimensions['A'].width = 52
ref.column_dimensions['B'].width = 16
wb.defined_names.add(DefinedName('film_list', attr_text="Справочник!$A$2:$A$%d" % ref.max_row))
ref.sheet_state = 'hidden'

def head(ws, titles):
    ws.append(titles)
    for i, _ in enumerate(titles, start=1):
        c = ws.cell(row=1, column=i)
        c.font = Font(bold=True)
        c.fill = HEAD_FILL
        c.alignment = Alignment(vertical='center')

def markup_sheet(name, rows, index):
    """Лист разметки: ссылка, выпадающий список, и рядом то, по чему человек решает."""
    ws = wb.create_sheet(name, index)
    # «Модель» — подсказка из очереди «Проверки» (пачка, 06.10); импорт читает только A–H
    head(ws, ['Обзор / видео', 'Фильм', 'Заголовок ролика', 'Канал', 'Дата', 'Проверено', 'Ошибка', 'Ещё фильмы', 'Модель'])
    for r, row in enumerate(rows, start=2):
        a = ws.cell(row=r, column=1, value=row['url'])
        a.hyperlink = row['url']
        a.font = LINK
        ws.cell(row=r, column=2, value=row['film'] or None)
        ws.cell(row=r, column=3, value=row['title'])
        ws.cell(row=r, column=4, value=row['channel'])
        ws.cell(row=r, column=5, value=row['date'])
        if row.get('checked'):
            ws.cell(row=r, column=6, value='да')      # решение человека — возвращаем как было
        if row.get('err'):
            ws.cell(row=r, column=7, value=row['err'])
        if row.get('also'):
            ws.cell(row=r, column=8, value=row['also'])
        if row.get('hint'):
            ws.cell(row=r, column=9, value=row['hint'])
    last = len(rows) + 1 + EMPTY   # пустые строки снизу — с тем же выпадающим списком

    # Выпадающий список. Предупреждение, а не запрет: если человек знает фильм, которого нет в
    # нашем списке, он должен суметь его вписать — такие строки нам и нужнее всего.
    dv = DataValidation(type='list', formula1='=film_list', allow_blank=True, showDropDown=False,
                        showInputMessage=True, showErrorMessage=True)
    dv.errorStyle = 'warning'
    dv.errorTitle = 'Нет в списке'
    dv.error = 'Такого фильма у нас нет. Оставьте как есть, если уверены, — разберём отдельно.'
    dv.promptTitle = 'Фильм'
    dv.prompt = 'Выберите из списка или впишите название.'
    dv.add('B2:B%d' % last)
    ws.add_data_validation(dv)

    dv2 = DataValidation(type='list', formula1='"да"', allow_blank=True, showDropDown=False,
                         showInputMessage=True, showErrorMessage=False)
    dv2.promptTitle = 'Проверено'
    dv2.prompt = 'Поставьте «да», когда посмотрели строку: колонка «Фильм» бывает заполнена догадкой машины.'
    dv2.add('F2:F%d' % last)
    ws.add_data_validation(dv2)

    # Вид ошибки опознавателя (02.10): по нему видно, какие правила чинить. «не фильм» и «не тот фильм»
    # снимают привязку и без «да»; «несколько фильмов» — фильм в колонке B верен, остальные — в H
    dv3 = DataValidation(type='list', formula1='"не тот фильм,не фильм,несколько фильмов,о франшизе,о человеке"', allow_blank=True,
                         showDropDown=False, showInputMessage=True, showErrorMessage=True)
    dv3.promptTitle = 'Ошибка'
    dv3.prompt = 'Если машина ошиблась: не тот фильм / не фильм вовсе / несколько фильмов / о франшизе (название — в «Фильм») / о человеке (имя — в «Фильм»).'
    dv3.add('G2:G%d' % last)
    ws.add_data_validation(dv3)
    for col, width in zip('ABCDEFGHI', (44, 46, 64, 22, 12, 12, 18, 40, 60)):
        ws.column_dimensions[col].width = width
    ws.freeze_panes = 'A2'
    ws.auto_filter.ref = 'A1:I%d' % (len(rows) + 1)
    return ws

markup_sheet('Обзоры', data['review'], 0)
markup_sheet('Эссе', data['essay'], 1)

# 2б. Корпус: самые обсуждаемые фильмы, у которых роликов мало, — их наполняют ссылками
corpus = data.get('corpus', [])
cw = max([len(c.get('links') or []) for c in corpus] + [1]) + 2
cs = wb.create_sheet('Корпус', 2)
head(cs, ['Фильм', 'Год', 'Назвали в постах', 'Роликов уже', 'из них обзоров', 'Ссылка на ролик'] + ['ещё ссылка'] * (cw - 1))
for r, c in enumerate(corpus, start=2):
    cs.cell(row=r, column=1, value=c['label'])
    cs.cell(row=r, column=2, value=c['year'] or None)
    cs.cell(row=r, column=3, value=c['talk'] or None)
    cs.cell(row=r, column=4, value=c['have'])
    cs.cell(row=r, column=5, value=c['reviews'] or None)
    for i, url in enumerate(c.get('links') or []):
        cs.cell(row=r, column=6 + i, value=url)
for i, width in enumerate((46, 8, 16, 12, 14) + (50,) * cw):
    cs.column_dimensions[get_column_letter(1 + i)].width = width
cs.freeze_panes = 'B2'
cs.auto_filter.ref = 'A1:%s%d' % (get_column_letter(5 + cw), len(corpus) + 1)

# 3. Обратный взгляд: фильмы, о которых у нас никто не высказался
gaps = wb.create_sheet('Без разбора', 2)
# ссылки — все, по одной в ячейке вправо от D (как их и вписывают); плюс одна пустая колонка для новой
links_of = lambda m: m.get('links') or ([m['link']] if m.get('link') else [])
wide = max([len(links_of(m)) for m in missing] + [1]) + 1
head(gaps, ['Фильм', 'Год', 'Назвали в постах', 'Знаю разбор — ссылка'] + ['ещё ссылка'] * (wide - 1))
for r, m in enumerate(missing, start=2):
    gaps.cell(row=r, column=1, value=m['label'])
    gaps.cell(row=r, column=2, value=m['year'] or None)
    gaps.cell(row=r, column=3, value=m['talk'] or None)
    for i, url in enumerate(links_of(m)):
        gaps.cell(row=r, column=4 + i, value=url)
for i, width in enumerate((46, 8, 18) + (50,) * wide):
    gaps.column_dimensions[get_column_letter(1 + i)].width = width
gaps.freeze_panes = 'A2'
gaps.auto_filter.ref = 'A1:%s%d' % (get_column_letter(3 + wide), len(missing) + 1)

# 4. Короткая памятка — в том же файле, чтобы не терялась
how = wb.create_sheet('Как размечать')
guessed = sum(1 for r in data['review'] + data['essay'] if r['film'] and not r.get('checked'))
for line in [
    'Что это: %d видеороликов с каналов о кино. Надо сказать, о каком фильме ролик.'
    % (len(data['review']) + len(data['essay'])),
    '',
    'Лист «Обзоры» (%d) — каналы, которые берут то, что смотрят все: BadComedian, КИНОКРИТИКА,'
    % len(data['review']),
    '  SokoL[off], Клим Жуков. Фильм у них почти всегда в заголовке — работа быстрая.',
    'Лист «Эссе» (%d) — каналы, которые берут то, о чём есть что сказать: КИНОЛИКБЕЗ, ЭПИЗОДЫ,'
    % len(data['essay']),
    '  Nuke, 4то за Персонаж. Заголовок у них бывает про что угодно — иногда надо открыть ролик.',
    '',
    'Это пачка: самые спорные и непроверенные строки, сверху — где машина сомневается сильнее',
    'всего. Колонка «Модель» — что думает языковая модель. Решённое уходит, при следующей',
    'синхронизации приходит новая пачка.',
    '',
    'Колонка «Фильм» — выпадающий список (стрелка справа в ячейке).',
    '  · %d строк уже заполнены догадкой машины — их надо подтвердить или поправить.' % guessed,
    '  · Ролик не о фильме (история, стендап, подборка) — выберите «%s».' % NOT_A_FILM,
    '  · Фильма нет в списке — впишите название руками, Excel только предупредит.',
    '  · «Проверено» = «да» — значит строку смотрел человек. Без этого догадку машины',
    '    не отличить от подтверждения.',
    '',
    'Быстрее всего начать с догадок: фильтр по колонке «Фильм» → «Непустые».',
    '',
    'Колонки «Заголовок ролика», «Канал», «Дата» — чтобы не открывать каждую ссылку.',
    'Ссылка в первой колонке кликается.',
    '',
    'Лист «Без разбора» (%d фильмов) — взгляд с другой стороны: наши фильмы, о которых'
    % len(missing),
    'не высказался ни один обзорщик и ни один эссеист. Сверху те, кого чаще называют в постах',
    'каналов, — это и есть очередь на разбор. Знаете разбор такого фильма — вставьте ссылку',
    'в последнюю колонку, она попадёт в разметку так же, как выбор из списка.',
    '',
    'Лист «Корпус» (%d фильмов) — самые обсуждаемые: из 300 фильмов, которые чаще всего называют'
    % len(corpus),
    'в постах авторов, те, у кого роликов пока не больше двух. Сверху — самые обсуждаемые. Знаете',
    'обзор или разбор — вставьте ссылку на ролик в колонку «Ссылка на ролик» (и правее, если',
    'их несколько): она попадёт в разметку, а фильм в листе останется со своими ссылками.',
    '',
    'Внизу листов разметки по %d пустых строк с тем же списком — для роликов, которых тут нет.' % EMPTY,
    'Список фильмов лежит на скрытом листе «Справочник», там же ключи. Лист не трогать.',
]:
    how.append([line])
how.column_dimensions['A'].width = 100

del wb['Sheet']
os.makedirs(os.path.dirname(OUT), exist_ok=True)
wb.save(OUT)
print('обзоры %d, эссе %d (с догадкой %d), корпус %d, без разбора %d, фильмов в списке %d'
      % (len(data['review']), len(data['essay']), guessed, len(corpus), len(missing), len(films) + 1))
print('→ %s  (%.1f МБ)' % (OUT, os.path.getsize(OUT) / 1048576))
