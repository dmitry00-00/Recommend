# -*- coding: utf-8 -*-
"""Прочитать размеченный film_reviews.xlsx и сложить решения людей в tools/markup-verdicts.json.

    python3 tools/import-markup.py [--in файл.xlsx | --url <ссылка на Google-таблицу>] [--dry]

С `--url` таблица берётся прямо из Google Таблиц в текущем виде — годится обычная ссылка из
адресной строки, она сама превращается в выгрузку .xlsx. Нужен доступ «по ссылке» хотя бы на
чтение: скачивание идёт без входа в аккаунт. Закрытый док так не прочитать — будет HTML
страницы входа вместо файла, и скрипт об этом скажет.

Нужен openpyxl (его же ставит deploy/markup-xlsx.command в .cache/venv).

Решением человека считается не всякая заполненная ячейка: 697 строк таблица отдаёт уже
заполненными догадкой машины, и принять их за разметку — значит записать догадку в правду.
Человеческое решение — это либо «Проверено» = «да», либо значение, отличное от догадки.

Ключ фильма берётся со скрытого листа «Справочник» того же файла, а не из наших моков: если
человек прислал файл, собранный неделю назад, правильный ключ лежит именно в нём.
Что вписано руками и в списке не нашлось — не теряем: это очередь на добавление фильмов.
"""
import json, os, re, sys
from datetime import date
from openpyxl import load_workbook

def arg(name, default):
    return sys.argv[sys.argv.index('--' + name) + 1] if '--' + name in sys.argv[:-1] else default

ROOT = os.path.dirname(os.path.dirname(os.path.abspath(__file__)))
SRC = arg('in', os.path.join(ROOT, '.cache/markup/film_reviews.xlsx'))
URL = arg('url', None)
# С чем сверять: с тем, что было в таблице, когда её отдали людям. Для локального файла это
# последняя сборка, для Google — снимок на момент заливки (.cache/markup/pushed.json, его
# пишет markup-google.py push). Сверка с текущей сборкой врёт: сторож снял догадку, а в
# таблице она осталась — и старая ошибка машины выглядит правкой человека (27.09: два
# ролика про «Однажды в Голливуде» так чуть не вернулись к «Однажды» 2007 года).
REF = arg('ref', os.path.join(ROOT, '.cache/markup/film-reviews.json'))
OUT = os.path.join(ROOT, 'tools/markup-verdicts.json')
DRY = '--dry' in sys.argv
# реестр в базе приложения (06.10, tools/registry-lib.mts): писать разметку из таблицы мимо базы нельзя
_envl = os.path.join(os.path.dirname(os.path.dirname(os.path.abspath(__file__))), '.env.local')
if os.environ.get('REGISTRY_MODE') == 'server' or (os.path.exists(_envl) and any(l.strip() == 'REGISTRY_MODE=server' for l in open(_envl, encoding='utf-8'))):
    sys.exit('реестр в базе приложения — импорт из таблицы разметки выключен (ссылки — форма Google, tools/inbox.mts)')
NOT_A_FILM = '— не про фильм —'
VIDEO = re.compile(r'(?:v=|youtu\.be/|/shorts/|/live/)([A-Za-z0-9_-]{11})')

if URL:
    # из ссылки на док достаём id и берём выгрузку всей книги: скрытый «Справочник» с ключами
    # в ней тоже есть, и дальше код не различает, откуда файл
    import subprocess
    m = re.search(r'/spreadsheets/d/([A-Za-z0-9_-]+)', URL)
    if not m:
        sys.exit('не похоже на ссылку Google Таблиц: %s' % URL)
    dl = 'https://docs.google.com/spreadsheets/d/%s/export?format=xlsx' % m.group(1)
    SRC = os.path.join(ROOT, '.cache/markup/google.xlsx')
    os.makedirs(os.path.dirname(SRC), exist_ok=True)
    # качаем curl'ом, а не из питона: у него системное хранилище корневых сертификатов, а у
    # свежего python.org-сборки своего нет, и urllib падает на проверке сертификата Google
    r = subprocess.run(['curl', '-fsSL', '--location-trusted', '-o', SRC, dl], capture_output=True)
    if r.returncode != 0:
        sys.exit('не скачалось (curl %d: %s). Проверьте доступ «по ссылке» у дока.'
                 % (r.returncode, r.stderr.decode()[:200].strip()))
    head = open(SRC, 'rb').read(4)
    if head[:2] != b'PK':
        sys.exit('вместо книги пришла страница — скорее всего, у дока нет доступа по ссылке.')
    print('скачал из Google Таблиц: %.1f КБ' % (os.path.getsize(SRC) / 1024))

if not os.path.exists(SRC):
    sys.exit('нет файла %s — соберите deploy/markup-xlsx.command и отдайте его людям' % SRC)

wb = load_workbook(SRC, data_only=True)
# «Разметка» — имя из первой сборки (26.09, до разделения на ярусы); читаем и его
SHEETS = [n for n in ('Обзоры', 'Эссе', 'Разметка') if n in wb.sheetnames]
if not SHEETS:
    sys.exit('в книге нет листов разметки, есть: %s' % ', '.join(wb.sheetnames))
keys = {}                                    # ярлык фильма → ключ
if 'Справочник' in wb.sheetnames:
    ref = wb['Справочник']
    for label, key in ref.iter_rows(min_row=2, max_col=2, values_only=True):
        if label and key:
            keys[str(label).strip()] = str(key).strip()

# чем была заполнена колонка «Фильм», когда таблицу отдавали: чтобы отличить подтверждение
# от нетронутой догадки. Файла нет — считаем решением только «Проверено» = «да».
guess = {}
# Нет снимка — не знаем, что в таблице было машинным. Тогда решением считаем только
# «Проверено» = «да»: лучше недобрать разметку, чем записать догадку в правду.
STRICT = not os.path.exists(REF)
if STRICT:
    print('нет %s — считаю только строки с «Проверено» = «да»' % os.path.relpath(REF, ROOT))
if os.path.exists(REF):
    ref_data = json.load(open(REF, encoding='utf-8'))
    for r in ref_data.get('review', []) + ref_data.get('essay', []) + ref_data.get('rows', []):
        m = VIDEO.search(r['url'])
        if m:
            guess[m.group(1)] = r['film']

# Названия, вписанные руками без года (28.09): их опознаёт tools/resolve-markup-films.mts по
# Wikidata с учётом даты ролика (ремейк и тёзка — частый случай) и пишет сюда ролик → ключ.
# Импорт применяет найденное, а всё неопознанное складывает в .cache/markup/typed.json — это
# вход опознавателя на следующем шаге.
RESOLVED = os.path.join(ROOT, 'tools', 'markup-resolved.json')
TYPED = os.path.join(ROOT, '.cache', 'markup', 'typed.json')
resolved = json.load(open(RESOLVED, encoding='utf-8')).get('videos', {}) if os.path.exists(RESOLVED) else {}
typed = []

was, prev = {}, {}
if os.path.exists(OUT):
    prev = json.load(open(OUT, encoding='utf-8'))
    was = prev.get('videos', {})

verdicts, stats = dict(was), {'подтвердили': 0, 'поправили': 0, 'не про фильм': 0, 'не тот фильм': 0, 'о франшизе': 0, 'о человеке': 0,
                              'нет у нас': 0, 'пропустили': 0, 'со стороны фильма': 0, 'из корпуса': 0,
                              'изменили прежнее': 0, 'без изменений': 0, 'опознали по названию': 0,
                              'опознали, на проверку': 0}
rows_read = 0

# guess — выбор опознавателя из тёзок: «да» в «Проверено» снимает пометку, это тоже изменение
SAME = ('key', 'why', 'film', 'guess', 'err', 'also', 'alsoFilms', 'aboutTitle')

# Вид ошибки (колонка «Ошибка», 02.10): те же три значения, что у пульта ссылок; хэштеги тоже понимаем
def error_of(text):
    t = str(text or '').strip().lower().replace('ё', 'е')
    if not t:
        return None
    if re.search(r'не\s*тот|#нетот', t):
        return 'не тот фильм'
    if re.search(r'не\s*фильм|не\s*про\s*фильм|#нефильм', t):
        return 'не фильм'
    if re.search(r'несколько|#несколько', t):
        return 'несколько фильмов'
    # о франшизе (цикле) и о человеке (режиссёре, писателе) — категории владельца 02.10
    if re.search(r'франшиз|вселенн|цикл|#франшиза', t):
        return 'о франшизе'
    if re.search(r'о\s*человеке|режиссер|писател|#очеловеке', t):
        return 'о человеке'
    return None

def also_of(text):
    """«Ещё фильмы»: ярлыки через «;» или с новой строки → ключи; неизвестные — названием"""
    ks, names = [], []
    for part in re.split(r'[;\n]+', str(text or '')):
        p = part.strip()
        if not p:
            continue
        (ks if p in keys else names).append(keys.get(p, p))
    return ks, names

# Вкладка «Проверка» пульта (tools/check-desk.mts, 02.10) пишет решения сразу сюда, как и правки
# мимо таблицы (пульт ссылок, разборы ошибок — `from`: desk, ops10, tv7…). Строка таблицы, отданной
# людям раньше, их не перетирает: решение не из таблицы, принятое в день заливки или позже, новее
# любой её ячейки (у ячеек нет своего времени — сравниваем с днём заливки). 06.10 защита была только
# у «Проверки», и круг markup-sync вернул 17 исправленных ошибок к прежним ячейкам таблицы.
FROM_SHEET = (None, 'gap', 'corpus')
PUSHED_DAY = ''
if os.path.exists(REF):
    PUSHED_DAY = date.fromtimestamp(os.path.getmtime(REF)).isoformat()
newer_in_desk = []

def put(vid, v, bucket):
    old = was.get(vid)
    if old and old.get('from') not in FROM_SHEET and old.get('at', '') >= PUSHED_DAY and v.get('from') in FROM_SHEET \
            and any(old.get(k) != v.get(k) for k in SAME):
        newer_in_desk.append(vid)
        return
    # таблица возвращается из Google с прежними решениями: их не считаем заново и дату не
    # трогаем, иначе каждый круг выглядел бы как сотни новых подтверждений
    if old and all(old.get(k) == v.get(k) for k in SAME):
        stats['без изменений'] += 1
        return
    v['at'] = date.today().isoformat()
    stats[bucket] += 1
    if old:
        stats['изменили прежнее'] += 1
    verdicts[vid] = v

for name in SHEETS:
    ws = wb[name]
    rows_read += ws.max_row - 1
    for row in ws.iter_rows(min_row=2, max_col=8, values_only=True):
        url, film, _title, _ch, _date, checked, err_cell, also_cell = (list(row) + [None] * 8)[:8]
        if not url:
            continue
        m = VIDEO.search(str(url))
        if not m:
            continue
        vid = m.group(1)
        # число в ячейке («1917») openpyxl отдаёт дробью
        film = str(int(film)) if isinstance(film, float) and film.is_integer() else (str(film).strip() if film else '')
        ok = str(checked).strip().lower() in ('да', 'yes', 'x', '+', 'true', '1') if checked else False
        before = guess.get(vid, '')
        err = error_of(err_cell)
        also_keys, also_names = also_of(also_cell)
        extra = {**({'err': err} if err else {}), **({'also': also_keys} if also_keys else {}),
                 **({'alsoFilms': also_names} if also_names else {})}
        # «не фильм» и «не тот фильм» — решение и без «да»: привязку снимаем; фильм в колонке — та
        # самая ошибочная догадка (или исправленный фильм, если его поменяли — тогда это поправка)
        if err == 'не фильм':
            put(vid, {'key': None, 'why': 'не про фильм', **extra}, 'не про фильм')
            continue
        # о франшизе или человеке: к фильму не привязываем; название цели — из колонки «Фильм», если
        # его там поменяли (нетронутая догадка — это название фильма, а не франшизы)
        if err in ('о франшизе', 'о человеке'):
            about = {'aboutKind': 'universe' if err == 'о франшизе' else 'person'}
            if film and film != before and film != NOT_A_FILM:
                about['aboutTitle'] = re.sub(r'\s*\((?:сериал,?\s*)?(?:19|20)\d{2}\)\s*$', '', film).strip()
            put(vid, {'key': None, 'why': err, **about, **extra}, err)
            continue
        if err == 'не тот фильм' and (not film or film == before):
            put(vid, {'key': None, 'why': 'не тот фильм', **({'film': film} if film else {}), **extra}, 'не тот фильм')
            continue
        if err:
            ok = True   # вид ошибки поставил человек — строку он смотрел

        if not film:
            stats['пропустили'] += 1
            continue
        if not ok and (STRICT or film == before):
            stats['пропустили'] += 1       # догадка, которую никто не смотрел, — не разметка
            continue

        if film == NOT_A_FILM:
            put(vid, {'key': None, 'why': 'не про фильм', **extra}, 'не про фильм')
        elif film in keys:
            put(vid, {'key': keys[film], 'film': film, **extra}, 'подтвердили' if film == before else 'поправили')
        elif resolved.get(vid, {}).get('typed') == film and resolved[vid].get('key'):
            r = resolved[vid]
            if r.get('sure'):
                put(vid, {'key': r['key'], 'film': r['label']}, 'опознали по названию')
            else:
                # выбор из тёзок по дате ролика — не решение человека: в таблицу уходит с годом и
                # пустым «Проверено», в индекс — как «не проверено»
                put(vid, {'key': r['key'], 'film': r['label'], 'guess': True}, 'опознали, на проверку')
        else:
            put(vid, {'key': None, 'why': 'нет у нас', 'film': film}, 'нет у нас')
            typed.append({'video': vid, 'film': film, 'title': str(_title or ''), 'sheet': name})

# Лист «Без разбора» — та же разметка, но со стороны фильма: строка знает фильм, человек
# приносит ссылку. Здесь подтверждать нечего, догадки машины в этом листе нет.
# Ссылок бывает много: владелец кладёт их в ряд — D, E, F… (29.09: 298 ссылок к 37 фильмам,
# до восьми на фильм), иногда по нескольку в одной ячейке. Ссылка на канал, а не на ролик,
# разметкой не считается — такие собираются в отчёт.
gap_channels = []
if 'Без разбора' in wb.sheetnames:
    for row in wb['Без разбора'].iter_rows(min_row=2, values_only=True):
        film, links = (row[0] if row else None), [c for c in row[3:] if c]
        if not links or not film:
            continue
        label = str(film).strip()
        if label not in keys:
            continue
        for cell in links:
            found = VIDEO.findall(str(cell))
            if not found and re.search(r'youtube\.com/(@|channel/|c/)', str(cell)):
                gap_channels.append('%s: %s' % (label, str(cell).split('?')[0]))
            for vid in found:
                put(vid, {'key': keys[label], 'film': label, 'from': 'gap'}, 'со стороны фильма')

# Лист «Корпус» — самые обсуждаемые фильмы, у которых роликов мало (01.10): тоже со стороны фильма,
# ссылки — от колонки F вправо
if 'Корпус' in wb.sheetnames:
    for row in wb['Корпус'].iter_rows(min_row=2, values_only=True):
        film, links = (row[0] if row else None), [c for c in row[5:] if c]
        if not links or not film:
            continue
        label = str(film).strip()
        if label not in keys:
            continue
        for cell in links:
            found = VIDEO.findall(str(cell))
            if not found and re.search(r'youtube\.com/(@|channel/|c/)', str(cell)):
                gap_channels.append('%s: %s' % (label, str(cell).split('?')[0]))
            for vid in found:
                put(vid, {'key': keys[label], 'film': label, 'from': 'corpus'}, 'из корпуса')

# Пульт ссылок (tools/links-desk.mts, 01.10): ролик к фильму, которого у нас нет, лежит в решениях как
# «нет у нас» с названием. В листах таблицы его нет — поэтому опознаватель (resolve-markup-films.mts)
# получает его отсюда, а найденное им применяется здесь же на втором круге импорта
for vid, v in list(verdicts.items()):
    if v.get('from') not in ('desk', 'check') or v.get('key') or v.get('why') != 'нет у нас' or not v.get('film'):
        continue
    r = resolved.get(vid, {})
    if r.get('typed') == v['film'] and r.get('key'):
        nv = {'key': r['key'], 'film': r['label'], 'from': 'desk'}
        if not r.get('sure'):
            nv['guess'] = True
        nv['from'] = v.get('from')
        put(vid, nv, 'опознали по названию' if r.get('sure') else 'опознали, на проверку')
    else:
        typed.append({'video': vid, 'film': v['film'], 'title': v.get('title', ''),
                      'sheet': 'проверка' if v.get('from') == 'check' else 'пульт ссылок'})

# прочие разделы файла (posts — решения по постам Telegram) импорт не трогает, но и не теряет:
# до 06.10 файл переписывался одним разделом videos, и решения по постам пропадали
body = {**prev,
        '//': 'Ручная разметка «ролик → фильм». Пишет tools/import-markup.py из film_reviews.xlsx,'
              ' читает tools/build-essay-index.mts. Правда сильнее догадки: перегенерация индекса'
              ' её не сотрёт. Править руками можно, ключ — id ролика на YouTube.',
        'updated': date.today().isoformat(),
        'videos': dict(sorted(verdicts.items()))}
os.makedirs(os.path.dirname(TYPED), exist_ok=True)
json.dump({'rows': typed}, open(TYPED, 'w', encoding='utf-8'), ensure_ascii=False, indent=1)
if not DRY:
    json.dump(body, open(OUT, 'w', encoding='utf-8'), ensure_ascii=False, indent=1, sort_keys=False)

print('прочитано строк: %d (листы: %s)' % (rows_read, ', '.join(SHEETS)))
for k, n in stats.items():
    print('  %s: %d' % (k, n))
print('решений всего в %s: %d%s' % (os.path.relpath(OUT, ROOT), len(verdicts), ' (--dry, файл не тронут)' if DRY else ''))
if newer_in_desk:
    print('решения не из таблицы («Проверка», пульт, исправления) новее неё — оставил их (%d): %s' % (len(newer_in_desk), ', '.join(newer_in_desk[:10])))
if gap_channels:
    print('в «Без разбора» ссылки на канал, а не на ролик (%d) — разметкой не считаю: %s' % (len(gap_channels), '; '.join(gap_channels)))
# «нет у нас» без названия бывает: фильм не опознан, но и не наш (исправления ошибок 06.10)
unknown = [v['film'] for v in verdicts.values() if v.get('why') == 'нет у нас' and v.get('film')]
if unknown:
    print('фильмов, которых у нас нет (%d): %s' % (len(unknown), ', '.join(sorted(set(unknown))[:10])))
